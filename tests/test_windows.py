"""Tests for the Windows engine. They run on any OS (the engine is plain PyTorch);
extraction tests need pypdf and python-docx, the model test needs the downloaded model."""
import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / "core"), str(ROOT / "windows")]
import engine_torch  # noqa: E402

HAS_DOC_LIBS = all(importlib.util.find_spec(m) for m in ("pypdf", "docx"))


def tiny_pdf(text):
    """A minimal one-page PDF showing `text` in Helvetica (fake data only)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    out, offsets = bytearray(b"%PDF-1.4\n"), []
    for i, body in enumerate(objs, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1)
    out += b"".join(b"%010d 00000 n \n" % o for o in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    return bytes(out)


class Extract(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())

    def test_text_utf8_and_cp1252(self):
        (self.dir / "a.txt").write_bytes("Café Jane".encode("utf-8"))
        (self.dir / "b.csv").write_bytes("Caf\xe9 Jane".encode("cp1252"))
        self.assertEqual(engine_torch.extract_text(self.dir / "a.txt"), "Café Jane")
        self.assertEqual(engine_torch.extract_text(self.dir / "b.csv"), "Café Jane")

    def test_html_drops_scripts(self):
        (self.dir / "a.html").write_text("<p>Jane Doe</p><script>var x=1;</script><p>Call 555</p>")
        text = engine_torch.extract_text(self.dir / "a.html")
        self.assertIn("Jane Doe", text)
        self.assertNotIn("var x", text)

    def test_unsupported_and_empty(self):
        (self.dir / "a.xyz").write_text("x")
        (self.dir / "b.txt").write_text("  \n ")
        with self.assertRaises(ValueError):
            engine_torch.extract_text(self.dir / "a.xyz")
        with self.assertRaises(ValueError):
            engine_torch.extract_text(self.dir / "b.txt")

    @unittest.skipUnless(HAS_DOC_LIBS, "pypdf and python-docx required")
    def test_pdf_and_docx(self):
        import docx
        (self.dir / "a.pdf").write_bytes(tiny_pdf("Statement for Jane Testperson"))
        d = docx.Document()
        d.add_paragraph("Statement for Jane Testperson")
        d.add_table(rows=1, cols=2).rows[0].cells[0].text = "Acct 12345678"
        d.save(self.dir / "a.docx")
        self.assertIn("Jane Testperson", engine_torch.extract_text(self.dir / "a.pdf"))
        docx_text = engine_torch.extract_text(self.dir / "a.docx")
        self.assertIn("Jane Testperson", docx_text)
        self.assertIn("Acct 12345678", docx_text)

    def test_output_path_never_overwrites(self):
        first = engine_torch.output_path(self.dir / "s.pdf", self.dir)
        first.write_text("x")
        self.assertEqual(first.name, "s (redacted).txt")
        self.assertEqual(engine_torch.output_path(self.dir / "s.pdf", self.dir).name, "s (redacted 2).txt")


def _torch_model_available():
    if importlib.util.find_spec("opf") is None or importlib.util.find_spec("mlx") is not None:
        return False  # the Mac environment has MLX instead; this test is for the PyTorch engine
    try:
        from huggingface_hub import snapshot_download
        snapshot_download(engine_torch.MODEL_REPO, revision=engine_torch.MODEL_REVISION,
                          allow_patterns=["original/*"], local_files_only=True)
        return True
    except Exception:
        return False


@unittest.skipUnless(_torch_model_available(), "PyTorch environment and model required")
class TorchModel(unittest.TestCase):
    def test_fake_record_is_redacted(self):
        from redact_core import redact_with
        text = "STATEMENT FOR JOHN QWERTYSON\nCall Jane Testperson at 555-123-4567.\n"
        out = redact_with(engine_torch.build_redactor(), text, ["Testperson"])
        for leaked in ("QWERTYSON", "Jane", "Testperson", "555-123-4567"):
            self.assertNotIn(leaked, out)


if __name__ == "__main__":
    unittest.main()
