"""localredact Windows engine: OpenAI Privacy Filter via its own PyTorch code.

Runs on an NVIDIA GPU when PyTorch sees one (CUDA), otherwise on the CPU (slower:
minutes for a long document). Also turns PDF, Word, HTML and text files into text.

CLI: python engine_torch.py INPUT OUTPUT.txt [ALWAYS_REDACT.txt]
     python engine_torch.py --download-only   (used by install.ps1)
"""
import os
import sys
from html.parser import HTMLParser
from pathlib import Path

HERE = Path(__file__).resolve().parent
# Installed: redact_core.py sits next to this file. In the repo: ../core.
sys.path[:0] = [str(HERE), str(HERE.parent / "core")]
from redact_core import read_terms, redact_with  # noqa: E402

MODEL_REPO = "openai/privacy-filter"
MODEL_REVISION = "7ffa9a043d54d1be65afb281eddf0ffbe629385b"
TEXT_EXTS = {".txt", ".md", ".csv", ".tsv", ".json", ".log", ".text"}
SUPPORTED_EXTS = TEXT_EXTS | {".pdf", ".docx", ".html", ".htm"}


def download_model():
    """Fetch the pinned checkpoint (~2.8 GB once) and return its directory."""
    from huggingface_hub import snapshot_download

    path = snapshot_download(MODEL_REPO, revision=MODEL_REVISION, allow_patterns=["original/*"])
    return str(Path(path) / "original")


def device_name():
    import torch

    return "cuda" if torch.cuda.is_available() else "cpu"


def build_redactor():
    # opf's GPU path defaults to Triton kernels, which have no official Windows build;
    # its plain-PyTorch path gives the same results.
    os.environ.setdefault("OPF_MOE_TRITON", "0")
    from opf._api import OPF

    return OPF(model=download_model(), device=device_name())


class _HTMLText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts, self._skip = [], 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style"):
            self._skip += 1
        elif tag in ("p", "br", "div", "li", "tr", "h1", "h2", "h3", "h4", "h5", "h6"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style") and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def _read_text_file(path):
    raw = Path(path).read_bytes()
    for enc in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    return raw.decode("utf-8", errors="replace")


def extract_text(path):
    """Plain text of a supported document. Raises ValueError for unsupported types or no text."""
    ext = Path(path).suffix.lower()
    if ext in TEXT_EXTS:
        text = _read_text_file(path)
    elif ext == ".pdf":
        from pypdf import PdfReader

        text = "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    elif ext == ".docx":
        import docx

        d = docx.Document(path)
        lines = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                lines.append("\t".join(cell.text for cell in row.cells))
        text = "\n".join(lines)
    elif ext in (".html", ".htm"):
        parser = _HTMLText()
        parser.feed(_read_text_file(path))
        text = "".join(parser.parts)
    else:
        raise ValueError(f"Unsupported file type: {ext or '(none)'} (use txt, md, csv, docx, html or pdf)")
    if not text.strip():
        raise ValueError(f"No text found in {Path(path).name} (a scanned PDF has images, not text).")
    return text


def output_path(src, out_dir):
    """'<name> (redacted).txt' in out_dir, numbered instead of overwriting."""
    base = Path(src).stem
    out = Path(out_dir) / f"{base} (redacted).txt"
    n = 2
    while out.exists():
        out = Path(out_dir) / f"{base} (redacted {n}).txt"
        n += 1
    return out


def main():
    if sys.argv[1:] == ["--download-only"]:
        print(download_model())
        return
    src, dst = sys.argv[1], sys.argv[2]
    # A list that exists but can't be read raises here: it must fail loudly, never be skipped.
    terms = read_terms(sys.argv[3]) if len(sys.argv) > 3 and os.path.exists(sys.argv[3]) else []
    Path(dst).write_text(redact_with(build_redactor(), extract_text(src), terms), encoding="utf-8")


if __name__ == "__main__":
    main()
